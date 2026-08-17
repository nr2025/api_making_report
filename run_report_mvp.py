from __future__ import annotations

import base64
import io
import json
import re
import sys
import urllib.error
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

try:
    import fitz
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

try:
    from PIL import Image
except ImportError:
    Image = None


ROOT_DIR = Path(__file__).resolve().parent
SETTINGS_PATH = ROOT_DIR / "настройки.yaml"
INPUT_DIR_DEFAULT = ROOT_DIR / "вход"
PROMPTS_DIR_DEFAULT = ROOT_DIR / "промпты"
OUTPUT_DIR_DEFAULT = ROOT_DIR / "выход"
TEMPLATE_PATH_DEFAULT = ROOT_DIR / "шаблон.docx"
SUPPORTED_EXTENSIONS = {".docx", ".doc", ".pdf", ".jpg", ".jpeg", ".png", ".txt"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png"}
PDF_TEXT_MIN_CHARS_PER_PAGE = 50
MAX_IMAGE_SIDE = 1536
ALLOWED_REASONING_EFFORTS = ("low", "medium", "xhigh")
DEFAULT_REASONING_EFFORT = "medium"
REASONING_MARKER_RE = re.compile(
    r"^REASONING:\s*(low|medium|xhigh)\s*$",
    re.IGNORECASE,
)


def проверить_окружение() -> None:
    """Жёсткая проверка обязательных библиотек до обработки документов."""
    checks: list[tuple[str, str, str]] = [
        ("fitz", "PyMuPDF", "pymupdf"),
        ("PIL", "Pillow", "Pillow"),
        ("docx", "python-docx", "python-docx"),
        ("yaml", "PyYAML", "PyYAML"),
        ("win32com.client", "pywin32", "pywin32"),
    ]
    missing_names: list[str] = []
    missing_packages: list[str] = []
    for module_name, display_name, pip_name in checks:
        try:
            __import__(module_name)
        except ImportError:
            missing_names.append(display_name)
            missing_packages.append(pip_name)
    if missing_names:
        packages = " ".join(missing_packages)
        raise RuntimeError(
            f"Отсутствуют библиотеки: {', '.join(missing_names)}. "
            f"Python: {sys.executable}. "
            f"Установите: {sys.executable} -m pip install {packages}"
        )

@dataclass
class Settings:
    api_url: str
    api_key: str
    model: str
    enable_thinking: bool | None
    repetition_penalty: float
    input_dir: Path
    prompts_dir: Path
    output_dir: Path
    template_path: Path
    num_predict: int
    request_timeout_sec: int
    max_json_retries: int
    reasoning_effort: str


@dataclass
class PreparedInput:
    text_documents: list[dict[str, str]]
    images_base64: list[tuple[str, str]]
    all_documents: list[str]
    unreadable_documents: list[str] = field(default_factory=list)


@dataclass
class PromptRunResult:
    prompt_file: Path
    output_json: Path
    parsed_json: dict | list | None
    raw_response_text: str


def parse_simple_yaml(path: Path) -> dict[str, str]:
    settings: dict[str, str] = {}
    raw_content = path.read_text(encoding="utf-8-sig")
    for raw_line in raw_content.splitlines():
        line = raw_line.strip()
        if not line or line.startswith("#") or ":" not in line:
            continue
        key, value = line.split(":", 1)
        settings[key.strip()] = value.strip().strip("'\"")
    return settings


def parse_optional_bool(value: str | None) -> bool | None:
    if value is None:
        return None
    normalized = value.strip().lower()
    if normalized in {"true", "1", "yes", "on"}:
        return True
    if normalized in {"false", "0", "no", "off"}:
        return False
    raise ValueError(f"Некорректное булево значение: {value}")


def parse_reasoning_effort(value: str | None, default: str = DEFAULT_REASONING_EFFORT) -> str:
    if value is None or not str(value).strip():
        return default
    normalized = str(value).strip().lower()
    if normalized not in ALLOWED_REASONING_EFFORTS:
        allowed = ", ".join(ALLOWED_REASONING_EFFORTS)
        raise ValueError(f"Некорректное reasoning_effort: {value}. Допустимо: {allowed}")
    return normalized


def extract_prompt_reasoning_effort(prompt_text: str, default: str) -> tuple[str, str]:
    """Если первая строка — REASONING: low|medium|xhigh, вернуть её значение и текст без маркера."""
    if not prompt_text:
        return default, prompt_text
    first_line, sep, rest = prompt_text.partition("\n")
    match = REASONING_MARKER_RE.match(first_line.strip())
    if not match:
        return default, prompt_text
    effort = parse_reasoning_effort(match.group(1), default)
    remainder = rest if sep else ""
    return effort, remainder.lstrip("\r\n")


def load_settings(path: Path) -> Settings:
    if not path.exists():
        raise FileNotFoundError(f"Не найден файл настроек: {path}")

    raw = parse_simple_yaml(path)
    api_url = raw.get("api_url")
    api_key = raw.get("api_key")
    model = raw.get("model")
    if not api_url or not api_key or not model:
        raise ValueError("В настройках должны быть указаны api_url, api_key и model")

    input_dir = Path(raw.get("input_dir", str(INPUT_DIR_DEFAULT)))
    prompts_dir = Path(raw.get("prompts_dir", str(PROMPTS_DIR_DEFAULT)))
    output_dir = Path(raw.get("output_dir", str(OUTPUT_DIR_DEFAULT)))
    template_path = Path(raw.get("template_path", str(TEMPLATE_PATH_DEFAULT)))
    enable_thinking = parse_optional_bool(raw.get("enable_thinking"))
    repetition_penalty = float(raw.get("repetition_penalty", "1.2"))
    num_predict = int(raw.get("num_predict", "16384"))
    request_timeout_sec = int(raw.get("request_timeout_sec", "300"))
    max_json_retries = int(raw.get("max_json_retries", "3"))
    reasoning_effort = parse_reasoning_effort(raw.get("reasoning_effort"), DEFAULT_REASONING_EFFORT)

    return Settings(
        api_url=api_url.replace(" ", "").rstrip("/"),
        api_key=api_key,
        model=model,
        enable_thinking=enable_thinking,
        repetition_penalty=repetition_penalty,
        input_dir=input_dir,
        prompts_dir=prompts_dir,
        output_dir=output_dir,
        template_path=template_path,
        num_predict=num_predict,
        request_timeout_sec=request_timeout_sec,
        max_json_retries=max(1, max_json_retries),
        reasoning_effort=reasoning_effort,
    )


def read_docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        xml_bytes = archive.read("word/document.xml")
    root = ET.fromstring(xml_bytes)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    chunks: list[str] = []
    for para in root.findall(".//w:p", ns):
        texts = [node.text for node in para.findall(".//w:t", ns) if node.text]
        paragraph = "".join(texts).strip()
        if paragraph:
            chunks.append(paragraph)
    return "\n".join(chunks)


def read_text_file(path: Path) -> str:
    encodings = ("utf-8", "utf-8-sig", "cp1251")
    for enc in encodings:
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def is_supported_document_file(file_path: Path) -> bool:
    return file_path.suffix.lower() in SUPPORTED_EXTENSIONS


def image_mime_type(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".png":
        return "image/png"
    if ext in {".jpg", ".jpeg"}:
        return "image/jpeg"
    return "application/octet-stream"


def image_output_format(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".png":
        return "PNG"
    return "JPEG"


def resize_image_bytes(image_bytes: bytes, output_format: str) -> bytes:
    if Image is None:
        raise RuntimeError("Для обработки изображений требуется Pillow (PIL).")

    with Image.open(io.BytesIO(image_bytes)) as img:
        width, height = img.size
        long_side = max(width, height)
        resized = img.copy()
        if long_side > MAX_IMAGE_SIDE:
            scale = MAX_IMAGE_SIDE / long_side
            new_size = (max(1, int(width * scale)), max(1, int(height * scale)))
            resampling = Image.Resampling.LANCZOS if hasattr(Image, "Resampling") else Image.LANCZOS
            resized = img.resize(new_size, resampling)

        output = io.BytesIO()
        fmt = output_format.upper()
        save_kwargs: dict = {}
        if fmt == "JPEG":
            if resized.mode not in {"RGB", "L"}:
                resized = resized.convert("RGB")
            save_kwargs = {"quality": 90, "optimize": True}
        resized.save(output, format=fmt, **save_kwargs)
        return output.getvalue()


def read_image_as_base64(path: Path) -> tuple[str, str]:
    mime = image_mime_type(path)
    output_format = image_output_format(path)
    resized_bytes = resize_image_bytes(path.read_bytes(), output_format)
    return mime, base64.b64encode(resized_bytes).decode("ascii")


def convert_doc_to_docx(path: Path) -> Path:
    try:
        import win32com.client
    except ImportError as exc:
        raise RuntimeError("Для обработки .doc требуется pywin32 (win32com).") from exc

    output_path = path.with_suffix(".converted.docx")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    doc_handle = None
    try:
        doc_handle = word.Documents.Open(str(path))
        doc_handle.SaveAs(str(output_path), FileFormat=16)
    finally:
        if doc_handle is not None:
            doc_handle.Close(False)
        word.Quit()
    return output_path


def read_pdf(path: Path) -> tuple[str, list[tuple[str, str]]]:
    if fitz is None:
        raise RuntimeError("Для обработки PDF требуется PyMuPDF (fitz).")

    text_parts: list[str] = []
    pages_text_len: list[int] = []
    image_payloads: list[tuple[str, str]] = []
    with fitz.open(path) as pdf:
        for page in pdf:
            page_text = page.get_text("text").strip()
            text_parts.append(page_text)
            pages_text_len.append(len(page_text))

        avg_page_chars = (sum(pages_text_len) / len(pages_text_len)) if pages_text_len else 0
        is_text_pdf = avg_page_chars >= PDF_TEXT_MIN_CHARS_PER_PAGE
        if is_text_pdf:
            full_text = "\n".join(part for part in text_parts if part).strip()
            return full_text, []

        for page in pdf:
            pix = page.get_pixmap(matrix=fitz.Matrix(2.0, 2.0))
            png_bytes = pix.tobytes("png")
            resized_png = resize_image_bytes(png_bytes, "PNG")
            image_payloads.append(("image/png", base64.b64encode(resized_png).decode("ascii")))

    return "", image_payloads


def read_current_subject_name(input_dir: Path) -> str:
    if not input_dir.exists():
        raise FileNotFoundError(f"Папка входа не найдена: {input_dir}")
    subject_path = input_dir / "субъект.txt"
    if not subject_path.exists():
        raise FileNotFoundError(f"Не найден файл с субъектом: {subject_path}")
    subject_name = read_text_file(subject_path).strip()
    if not subject_name:
        raise ValueError(f"Файл субъекта пуст: {subject_path}")
    return subject_name


def resolve_subject_folder(input_dir: Path, subject_name: str) -> Path:
    subject_dir = input_dir / subject_name
    if subject_dir.exists() and subject_dir.is_dir():
        return subject_dir

    known_subject_dirs = sorted(p.name for p in input_dir.iterdir() if p.is_dir())
    raise FileNotFoundError(
        f"Не найдена папка субъекта '{subject_name}' в {input_dir}. "
        f"Найденные подпапки: {known_subject_dirs}"
    )


def collect_documents(subject_dir: Path) -> PreparedInput:
    text_docs: list[dict[str, str]] = []
    images_base64: list[tuple[str, str]] = []
    readable_docs: list[str] = []
    unreadable_docs: list[str] = []

    for file_path in sorted(p for p in subject_dir.rglob("*") if p.is_file()):
        if file_path.name.lower() == "субъект.txt":
            continue
        if not is_supported_document_file(file_path):
            continue

        rel_path = str(file_path.relative_to(subject_dir))
        ext = file_path.suffix.lower()
        got_text = False
        got_images = False
        try:
            if ext == ".txt":
                content = read_text_file(file_path).strip()
                if content:
                    text_docs.append({"path": rel_path, "content": content})
                    got_text = True
            elif ext == ".docx":
                content = read_docx_text(file_path).strip()
                if content:
                    text_docs.append({"path": rel_path, "content": content})
                    got_text = True
            elif ext == ".doc":
                converted = convert_doc_to_docx(file_path)
                content = read_docx_text(converted).strip()
                if content:
                    text_docs.append({"path": rel_path, "content": content})
                    got_text = True
            elif ext == ".pdf":
                pdf_text, pdf_images = read_pdf(file_path)
                if pdf_text:
                    text_docs.append({"path": rel_path, "content": pdf_text})
                    got_text = True
                if pdf_images:
                    images_base64.extend(pdf_images)
                    got_images = True
            elif ext in IMAGE_EXTENSIONS:
                images_base64.append(read_image_as_base64(file_path))
                got_images = True
        except Exception:  # noqa: BLE001 — файл помечается как нечитаемый
            got_text = False
            got_images = False

        if got_text or got_images:
            readable_docs.append(rel_path)
        else:
            unreadable_docs.append(rel_path)

    return PreparedInput(
        text_documents=text_docs,
        images_base64=images_base64,
        all_documents=readable_docs,
        unreadable_documents=unreadable_docs,
    )


def prompt_sort_key(path: Path) -> tuple[int, str]:
    match = re.match(r"^(\d+)", path.stem)
    if match:
        return (int(match.group(1)), path.name)
    return (10**9, path.name)


def list_prompt_files(prompts_dir: Path) -> list[Path]:
    if not prompts_dir.exists():
        raise FileNotFoundError(f"Папка промптов не найдена: {prompts_dir}")
    files = sorted((p for p in prompts_dir.rglob("*.txt") if p.is_file()), key=prompt_sort_key)
    if not files:
        raise FileNotFoundError(f"В папке промптов нет .txt файлов: {prompts_dir}")
    return files


def subject_prefix(subject_name: str) -> str:
    return (
        f"Субъект исследования: {subject_name}.\n"
        "Если в задании не указано иное, данные извлекаются в отношении Субъекта.\n"
        "Рассуждай и отвечай только на русском языке. "
        'Ключ в JSON — строго "данные" русскими буквами.'
    )


def build_prompt_with_documents(
    subject_name: str,
    base_prompt: str,
    documents: Iterable[dict[str, str]],
    image_count: int,
) -> str:
    parts = [subject_prefix(subject_name), "", base_prompt.strip(), "", "Документы для анализа:"]
    for idx, doc in enumerate(documents, start=1):
        parts.append("")
        parts.append(f"=== Документ {idx}: {doc['path']} ===")
        parts.append(doc["content"])
    if image_count:
        parts.append("")
        parts.append(f"Приложено изображений для анализа: {image_count}")
    return "\n".join(parts).strip()


def build_prompt_with_previous_steps(
    subject_name: str,
    base_prompt: str,
    previous_steps: list[PromptRunResult],
) -> str:
    parts = [subject_prefix(subject_name), "", base_prompt.strip(), "", "Результаты предыдущих шагов:"]
    for step in previous_steps:
        parts.append("")
        parts.append(f"=== {step.prompt_file.stem} ===")
        if step.parsed_json is not None:
            parts.append(json.dumps(step.parsed_json, ensure_ascii=False, indent=2))
            continue

        if step.output_json.exists():
            try:
                disk_payload = json.loads(step.output_json.read_text(encoding="utf-8"))
                parts.append(json.dumps(disk_payload, ensure_ascii=False, indent=2))
                continue
            except (OSError, json.JSONDecodeError):
                pass

        parts.append(step.raw_response_text)
    return "\n".join(parts).strip()


def normalize_chat_completions_url(api_url: str) -> str:
    if re.search(r"/v1/chat/completions/?$", api_url):
        return api_url
    return f"{api_url}/v1/chat/completions"


def build_chat_messages(prompt: str, images_base64: list[tuple[str, str]] | None = None) -> list[dict]:
    if not images_base64:
        return [{"role": "user", "content": prompt}]

    content_items: list[dict] = [{"type": "text", "text": prompt}]
    for mime, encoded_data in images_base64:
        content_items.append(
            {
                "type": "image_url",
                "image_url": {"url": f"data:{mime};base64,{encoded_data}"},
            }
        )
    return [{"role": "user", "content": content_items}]


def _message_text_field(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        parts: list[str] = []
        for item in value:
            if isinstance(item, dict) and item.get("type") == "text":
                parts.append(str(item.get("text") or ""))
            elif isinstance(item, str):
                parts.append(item)
        return "\n".join(part for part in parts if part)
    return str(value)


def extract_chat_content(response_payload: dict) -> str:
    choices = response_payload.get("choices", [])
    if not choices:
        return ""
    message = choices[0].get("message", {})
    return _message_text_field(message.get("content"))


def extract_reasoning_text(response_payload: dict) -> str:
    """Текст рассуждений из ответа API. На разбор JSON и сборку Word не влияет."""
    if not isinstance(response_payload, dict):
        return ""
    choices = response_payload.get("choices") or []
    message = choices[0].get("message", {}) if choices else {}
    if not isinstance(message, dict):
        message = {}

    candidates: list[str] = []
    for source in (message, choices[0] if choices else {}, response_payload):
        if not isinstance(source, dict):
            continue
        for key in ("reasoning", "reasoning_content", "reasoning_text", "thinking"):
            text = _message_text_field(source.get(key)).strip()
            if text:
                candidates.append(text)

    content = _message_text_field(message.get("content"))
    thinking, _answer = split_thinking_and_answer(content)
    thinking = thinking.strip()
    if thinking:
        candidates.append(thinking)

    seen: set[str] = set()
    unique: list[str] = []
    for text in candidates:
        if text not in seen:
            seen.add(text)
            unique.append(text)
    return "\n\n".join(unique)


def call_ollama(
    api_url: str,
    api_key: str,
    model: str,
    enable_thinking: bool | None,
    repetition_penalty: float,
    prompt: str,
    num_predict: int,
    timeout_sec: int,
    images_base64: list[tuple[str, str]] | None = None,
    reasoning_effort: str = DEFAULT_REASONING_EFFORT,
) -> dict:
    url = normalize_chat_completions_url(api_url)
    payload = {
        "model": model,
        "messages": build_chat_messages(prompt, images_base64),
        "max_tokens": num_predict,
        "temperature": 0.6,
        "repetition_penalty": repetition_penalty,
        "reasoning_effort": reasoning_effort,
    }
    if enable_thinking is not None:
        payload["extra_body"] = {
            "chat_template_kwargs": {"enable_thinking": enable_thinking}
        }

    data = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    request = urllib.request.Request(
        url=url,
        data=data,
        headers={
            "Content-Type": "application/json",
            "Authorization": f"Bearer {api_key}",
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(request, timeout=timeout_sec) as response:
            body = response.read().decode("utf-8")
    except urllib.error.HTTPError as exc:
        details = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"vLLM HTTP {exc.code}: {details}") from exc
    except urllib.error.URLError as exc:
        raise RuntimeError(f"Не удалось подключиться к vLLM API: {exc}") from exc

    return json.loads(body)


THINK_END_MARKER = "</think>"
MARKDOWN_FENCE_RE = re.compile(r"```(?:json)?\s*\n?(.*?)```", re.DOTALL | re.IGNORECASE)
# Одна и та же подстрока подряд 30+ раз (не суммарные вхождения по тексту).
CONSECUTIVE_REPEAT_RE = re.compile(r"(.{1,64})\1{29,}", re.DOTALL)


def split_thinking_and_answer(raw_text: str) -> tuple[str, str]:
    """Шаг 1: отрезать всё до </think> включительно. Возвращает (рассуждения, ответ)."""
    marker_pos = raw_text.find(THINK_END_MARKER)
    if marker_pos == -1:
        return "", raw_text
    thinking = raw_text[:marker_pos]
    answer = raw_text[marker_pos + len(THINK_END_MARKER) :].strip()
    return thinking, answer


def strip_markdown_fences(text: str) -> str:
    """Шаг 2: снять markdown-обёртки; при нескольких блоках взять последний."""
    matches = MARKDOWN_FENCE_RE.findall(text)
    if not matches:
        return text
    return matches[-1].strip()


def extract_json_from_text(text: str) -> dict | list | None:
    """Шаг 3: распарсить JSON целиком, иначе фрагмент от первой { до последней }."""
    text = text.strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass

    start = text.find("{")
    end = text.rfind("}")
    if start == -1 or end == -1 or end < start:
        return None
    try:
        return json.loads(text[start : end + 1])
    except json.JSONDecodeError:
        return None


def normalize_payload_with_data_array(payload: dict) -> dict | None:
    """Шаг 4: принять только объект с ключом данные/data-массивом."""
    for key in ("данные", "data"):
        value = payload.get(key)
        if not isinstance(value, list):
            continue
        if key == "данные":
            return payload
        normalized_payload = dict(payload)
        normalized_payload["данные"] = value
        return normalized_payload
    return None


def try_parse_valid_payload(text: str) -> dict | None:
    """Шаги 2–4: markdown → JSON → проверка ключа данные/data."""
    if not text or not text.strip():
        return None
    candidate = strip_markdown_fences(text)
    parsed = extract_json_from_text(candidate)
    if not isinstance(parsed, dict):
        return None
    return normalize_payload_with_data_array(parsed)


def has_consecutive_repeated_substring(
    text: str,
    max_allowed_repeats: int = 30,
) -> bool:
    """Детектор дегенерации: только подряд идущие повторы одной подстроки."""
    if not text or max_allowed_repeats < 2:
        return False
    if max_allowed_repeats == 30:
        return CONSECUTIVE_REPEAT_RE.search(text) is not None
    pattern = re.compile(rf"(.{{1,64}})\1{{{max_allowed_repeats - 1},}}", re.DOTALL)
    return pattern.search(text) is not None


def call_until_valid_json(
    settings: Settings,
    prompt: str,
    images_base64: list[tuple[str, str]] | None = None,
    reasoning_effort: str | None = None,
) -> tuple[dict, dict | list | None, str | None]:
    last_response: dict = {}
    last_raw_invalid_response = ""
    retries = min(settings.max_json_retries, 3)
    effort = reasoning_effort or settings.reasoning_effort
    for _ in range(retries):
        last_response = call_ollama(
            api_url=settings.api_url,
            api_key=settings.api_key,
            model=settings.model,
            enable_thinking=settings.enable_thinking,
            repetition_penalty=settings.repetition_penalty,
            prompt=prompt,
            num_predict=settings.num_predict,
            timeout_sec=settings.request_timeout_sec,
            images_base64=images_base64,
            reasoning_effort=effort,
        )
        raw_text = extract_chat_content(last_response)
        thinking_text, answer_text = split_thinking_and_answer(raw_text)

        # Шаги 1–4 по ответу после </think> (или по всему тексту, если тега нет).
        payload = try_parse_valid_payload(answer_text)
        if payload is not None:
            return last_response, payload, None

        # Шаг 6: спасение из рассуждений (те же шаги 2–4), до перезапроса.
        # Детектор повторов к рассуждениям не применяется.
        salvaged = try_parse_valid_payload(thinking_text)
        if salvaged is not None:
            return last_response, salvaged, None

        # Шаг 5: дегенерация только по остатку после обрезки и только если JSON не принят.
        # Считаем исключительно подряд идущие повторы; к рассуждениям не применяем.
        # Затем перезапрос (и при дегенерации, и при обычной ошибке формата).
        last_raw_invalid_response = raw_text
        has_consecutive_repeated_substring(answer_text)
        continue
    return last_response, None, last_raw_invalid_response or extract_chat_content(last_response)


def ensure_dict_payload(value: dict | list | None) -> dict:
    if isinstance(value, dict):
        return value
    if isinstance(value, list):
        return {"данные": value}
    return {"данные": []}


def save_prompt_result_json(
    intermediate_dir: Path,
    prompt_file: Path,
    subject_name: str,
    docs: PreparedInput,
    ollama_response: dict,
    parsed_json: dict | list | None,
    format_error_raw_response: str | None = None,
    reasoning_effort: str | None = None,
) -> PromptRunResult:
    intermediate_dir.mkdir(parents=True, exist_ok=True)
    output_json = intermediate_dir / f"{prompt_file.stem}.json"
    raw_response_text = extract_chat_content(ollama_response)
    reasoning_text = extract_reasoning_text(ollama_response)
    has_format_error = format_error_raw_response is not None and parsed_json is None
    payload = {
        "meta": {
            "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "subject_name": subject_name,
            "prompt_file": str(prompt_file),
            "documents_count": len(docs.all_documents),
            "text_documents_count": len(docs.text_documents),
            "images_count": len(docs.images_base64),
            "documents": docs.all_documents,
            "unreadable_documents": list(docs.unreadable_documents),
            "format_error": has_format_error,
            "reasoning_effort": reasoning_effort,
        },
        "данные": ensure_dict_payload(parsed_json).get("данные", []),
        "parsed_json_response": parsed_json,
        "raw_response_text": raw_response_text,
        "reasoning_text": reasoning_text,
    }
    output_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

    if format_error_raw_response:
        error_path = intermediate_dir / f"{prompt_file.stem}.ошибка_формата.txt"
        error_body = (
            "ошибка формата\n"
            f"prompt: {prompt_file.name}\n"
            "Требуется JSON-объект с ключом \"данные\".\n\n"
            "Сырой ответ модели:\n"
            f"{format_error_raw_response}"
        )
        error_path.write_text(error_body, encoding="utf-8")

    return PromptRunResult(
        prompt_file=prompt_file,
        output_json=output_json,
        parsed_json=parsed_json,
        raw_response_text=raw_response_text,
    )


def set_paragraph_text_with_breaks(paragraph, text: str) -> None:
    paragraph.clear()
    parts = str(text).split("\n")
    run = paragraph.add_run(parts[0] if parts else "")
    for chunk in parts[1:]:
        run.add_break()
        run = paragraph.add_run(chunk)


def set_cell_text_with_breaks(cell, text: str) -> None:
    if not cell.paragraphs:
        paragraph = cell.add_paragraph("")
    else:
        paragraph = cell.paragraphs[0]
    set_paragraph_text_with_breaks(paragraph, text)
    for extra in cell.paragraphs[1:]:
        extra.clear()


def replace_placeholder_with_lines(paragraph, placeholder: str, lines: list[str]) -> bool:
    if placeholder not in paragraph.text:
        return False
    safe_lines = lines or ["данные не обнаружены"]
    prefix, suffix = paragraph.text.split(placeholder, 1)
    first_line = f"{prefix}{safe_lines[0]}{suffix}"
    set_paragraph_text_with_breaks(paragraph, first_line)

    anchor = paragraph
    for extra_line in safe_lines[1:]:
        new_paragraph = paragraph.insert_paragraph_before("")
        anchor._p.addnext(new_paragraph._p)  # noqa: SLF001
        set_paragraph_text_with_breaks(new_paragraph, extra_line)
        anchor = new_paragraph
    return True


def stringify_cell_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, (int, float, bool)):
        return str(value)
    return json.dumps(value, ensure_ascii=False)


def format_field_name(key: str) -> str:
    name = str(key).replace("_", " ").strip()
    if not name:
        return name
    return name[0].upper() + name[1:]


def object_to_text_lines(item: dict) -> str:
    parts: list[str] = []
    for key, value in item.items():
        text = stringify_cell_value(value).strip()
        if not text:
            continue
        parts.append(f"{format_field_name(key)}: {text}")
    return "\n".join(parts)


def find_placeholder_row(table, placeholder: str):
    for row in table.rows:
        for cell in row.cells:
            if placeholder in cell.text:
                return row
    return None


def is_data_table_for(table, placeholder: str) -> bool:
    """Таблица данных: метка во 2-й строке, над ней строка заголовков из непустых ячеек."""
    placeholder_row_index = None
    for index, row in enumerate(table.rows):
        if any(placeholder in cell.text for cell in row.cells):
            placeholder_row_index = index
            break
    if placeholder_row_index is None:
        return False
    # Первая строка таблицы должна быть непосредственно над строкой с меткой.
    if placeholder_row_index != 1:
        return False

    header = table.rows[0]
    if any(placeholder in cell.text for cell in header.cells):
        return False
    if not header.cells:
        return False
    if not all(cell.text.strip() for cell in header.cells):
        return False
    return True


def fill_table_for_objects(table, placeholder: str, values: list[dict]) -> bool:
    row = find_placeholder_row(table, placeholder)
    if row is None:
        return False

    table._tbl.remove(row._tr)  # noqa: SLF001
    rows_to_add = values if values else [{"данные": "данные не обнаружены"}]
    for item in rows_to_add:
        new_row = table.add_row()
        fields = list(item.items())
        n_cols = len(new_row.cells)
        if n_cols == 0:
            continue

        for idx in range(min(len(fields), n_cols)):
            set_cell_text_with_breaks(new_row.cells[idx], stringify_cell_value(fields[idx][1]))

        if len(fields) > n_cols:
            last_parts = [stringify_cell_value(fields[n_cols - 1][1])]
            for key, value in fields[n_cols:]:
                last_parts.append(f"{format_field_name(key)}: {stringify_cell_value(value)}")
            set_cell_text_with_breaks(new_row.cells[n_cols - 1], "\n".join(last_parts))
    return True


def replace_placeholder_in_paragraphs(document, placeholder: str, lines: list[str]) -> bool:
    changed = False
    for paragraph in document.paragraphs:
        if replace_placeholder_with_lines(paragraph, placeholder, lines):
            changed = True
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    if replace_placeholder_with_lines(paragraph, placeholder, lines):
                        changed = True
    return changed


def replace_text_in_paragraph(paragraph, old: str, new: str) -> None:
    if old not in paragraph.text:
        return
    for run in paragraph.runs:
        if old in run.text:
            run.text = run.text.replace(old, new)
    # Если «g.» разбит по нескольким run — пересобираем текст абзаца целиком.
    if old in paragraph.text:
        set_paragraph_text_with_breaks(paragraph, paragraph.text.replace(old, new))


def apply_document_text_fixes(document) -> None:
    """Автозамены в собранном Word-документе."""
    for paragraph in document.paragraphs:
        replace_text_in_paragraph(paragraph, "g.", "г.")
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_text_in_paragraph(paragraph, "g.", "г.")


def insert_paragraph_after(paragraph, text: str, bold: bool = False):
    new_p = docx.oxml.OxmlElement("w:p")
    paragraph._p.addnext(new_p)  # noqa: SLF001
    new_para = docx.text.paragraph.Paragraph(new_p, paragraph._parent)
    run = new_para.add_run(text)
    run.bold = bold
    return new_para


def insert_unreadable_warning(document, unreadable_documents: list[str]) -> None:
    if not unreadable_documents:
        return
    names = ", ".join(unreadable_documents)
    warning = (
        "⚠ ВНИМАНИЕ: следующие документы не были обработаны: "
        f"{names}. Отчёт составлен без их учёта."
    )
    if not document.paragraphs:
        paragraph = document.add_paragraph("")
        run = paragraph.add_run(warning)
        run.bold = True
        return
    insert_paragraph_after(document.paragraphs[0], warning, bold=True)


def build_word_report(
    template_path: Path,
    output_docx: Path,
    prompt_results: list[PromptRunResult],
    unreadable_documents: list[str] | None = None,
) -> Path:
    if docx is None:
        raise RuntimeError("Для сборки Word нужен пакет python-docx.")
    if not template_path.exists():
        raise FileNotFoundError(f"Не найден шаблон Word: {template_path}")

    document = docx.Document(str(template_path))
    insert_unreadable_warning(document, unreadable_documents or [])
    payload_by_tag: dict[str, dict] = {}
    for step in prompt_results:
        try:
            payload_by_tag[step.prompt_file.stem] = json.loads(step.output_json.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            payload_by_tag[step.prompt_file.stem] = {"данные": []}

    for tag, payload in payload_by_tag.items():
        placeholder = f"{{{{{tag}}}}}"
        meta = payload.get("meta", {})
        has_format_error = bool(meta.get("format_error"))
        raw_values = payload.get("данные", [])
        values = raw_values if isinstance(raw_values, list) else []
        if has_format_error:
            replace_placeholder_in_paragraphs(
                document,
                placeholder,
                ["⚠ Ошибка обработки раздела, требуется ручная проверка"],
            )
            continue
        is_object_values = values and all(isinstance(item, dict) for item in values)

        if is_object_values:
            data_table = None
            for table in document.tables:
                if is_data_table_for(table, placeholder):
                    data_table = table
                    break
            if data_table is not None:
                fill_table_for_objects(data_table, placeholder, values)
            else:
                text_blocks = [object_to_text_lines(item) for item in values]
                text_blocks = [block for block in text_blocks if block] or ["данные не обнаружены"]
                replace_placeholder_in_paragraphs(document, placeholder, text_blocks)
            continue

        string_lines = [str(item) for item in values if isinstance(item, str)]
        if not string_lines:
            string_lines = ["данные не обнаружены"]
        replace_placeholder_in_paragraphs(document, placeholder, string_lines)

    apply_document_text_fixes(document)

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    try:
        document.save(str(output_docx))
    except PermissionError as exc:
        lock_file = output_docx.parent / f"~${output_docx.name}"
        raise RuntimeError(
            "Не удалось сохранить итоговый Word-файл. "
            f"Путь: {output_docx}. "
            "Скорее всего файл открыт в Word или заблокирован другим процессом. "
            "Закройте документ, удалите временный lock-файл (если есть) "
            f"{lock_file} и запустите скрипт снова."
        ) from exc
    return output_docx


def process_subject_folder(
    settings: Settings,
    subject_dir: Path,
    subject_name: str,
    progress_callback=None,
) -> tuple[Path, Path, list[str]]:
    проверить_окружение()
    prepared = collect_documents(subject_dir)
    if not prepared.text_documents and not prepared.images_base64:
        raise RuntimeError("Ни один документ не удалось прочитать")

    prompt_files = list_prompt_files(settings.prompts_dir)
    output_subject_dir = settings.output_dir / subject_dir.name
    intermediate_dir = output_subject_dir / "промежуточные"
    prompt_runs: list[PromptRunResult] = []
    total_steps = len(prompt_files)

    for index, prompt_file in enumerate(prompt_files):
        if progress_callback is not None:
            progress_callback(index + 1, total_steps)
        prompt_text = read_text_file(prompt_file)
        reasoning_effort, prompt_text = extract_prompt_reasoning_effort(
            prompt_text, settings.reasoning_effort
        )
        is_last = index == len(prompt_files) - 1
        if is_last:
            full_prompt = build_prompt_with_previous_steps(subject_name, prompt_text, prompt_runs)
            images = None
        else:
            full_prompt = build_prompt_with_documents(
                subject_name=subject_name,
                base_prompt=prompt_text,
                documents=prepared.text_documents,
                image_count=len(prepared.images_base64),
            )
            images = prepared.images_base64

        response, parsed_json, format_error_raw_response = call_until_valid_json(
            settings,
            full_prompt,
            images,
            reasoning_effort=reasoning_effort,
        )
        prompt_runs.append(
            save_prompt_result_json(
                intermediate_dir=intermediate_dir,
                prompt_file=prompt_file,
                subject_name=subject_name,
                docs=prepared,
                ollama_response=response,
                parsed_json=parsed_json,
                format_error_raw_response=format_error_raw_response,
                reasoning_effort=reasoning_effort,
            )
        )

    report_docx_path = output_subject_dir / f"{subject_dir.name}.docx"
    build_word_report(
        settings.template_path,
        report_docx_path,
        prompt_runs,
        unreadable_documents=prepared.unreadable_documents,
    )
    return intermediate_dir, report_docx_path, list(prepared.unreadable_documents)


def main() -> None:
    settings = load_settings(SETTINGS_PATH)
    subject_name = read_current_subject_name(settings.input_dir)
    subject_dir = resolve_subject_folder(settings.input_dir, subject_name)
    intermediate_dir, report_docx_path, _unreadable = process_subject_folder(
        settings, subject_dir, subject_name
    )
    print(f"Промежуточные JSON: {intermediate_dir}")
    print(f"Итоговый Word: {report_docx_path}")


if __name__ == "__main__":
    main()
